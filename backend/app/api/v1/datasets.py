import csv
import io
from pathlib import Path
from typing import Any

from fastapi import APIRouter, Depends, File, HTTPException, UploadFile, status

from app.api.deps import current_user_id
from app.core.config import settings
from app.repositories.in_memory import repository
from app.schemas.dataset import DatasetUploadResponse


router = APIRouter(prefix="/datasets", tags=["datasets"])


def _normalize_columns(raw_columns: list[Any]) -> list[str]:
    columns = [str(column).strip() if column is not None else "" for column in raw_columns]
    if not columns or any(not column for column in columns):
        raise HTTPException(
            status_code=400,
            detail={"code": "invalid_columns", "message": "Column names must not be empty"},
        )
    if len(set(columns)) != len(columns):
        raise HTTPException(
            status_code=400,
            detail={"code": "duplicate_columns", "message": "Column names must be unique after normalization"},
        )
    return columns


def _parse_csv(content: bytes) -> tuple[list[str], list[dict[str, Any]]]:
    try:
        decoded = content.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise HTTPException(
            status_code=400,
            detail={"code": "invalid_encoding", "message": "CSV must be UTF-8"},
        ) from exc
    try:
        csv.field_size_limit(settings.max_csv_field_length)
        reader = csv.DictReader(io.StringIO(decoded))
        raw_columns = reader.fieldnames or []
        columns = _normalize_columns(raw_columns)
        rows: list[dict[str, Any]] = []
        for row_number, row in enumerate(reader, start=2):
            if row_number - 1 > settings.max_dataset_rows:
                raise HTTPException(
                    status_code=413,
                    detail={"code": "dataset_too_large", "message": f"Dataset exceeds the {settings.max_dataset_rows}-row MVP limit"},
                )
            if None in row:
                raise HTTPException(
                    status_code=400,
                    detail={"code": "extra_csv_fields", "message": f"CSV row {row_number} has more fields than its header"},
                )
            rows.append(
                {column: row.get(raw_column) for raw_column, column in zip(raw_columns, columns)}
            )
        return columns, rows
    except csv.Error as exc:
        raise HTTPException(
            status_code=400,
            detail={"code": "invalid_csv", "message": "CSV contains a malformed or oversized field"},
        ) from exc


def _parse_xlsx(content: bytes) -> tuple[list[str], list[dict[str, Any]]]:
    try:
        from datetime import date, datetime

        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - dependency is part of runtime requirements
        raise HTTPException(
            status_code=503,
            detail={"code": "xlsx_unavailable", "message": "XLSX parser is not installed"},
        ) from exc
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        worksheet = workbook.active
        row_iterator = worksheet.iter_rows(values_only=True)
        raw_columns = list(next(row_iterator))
        columns = _normalize_columns(raw_columns)
        rows: list[dict[str, Any]] = []
        for row_number, values in enumerate(row_iterator, start=2):
            if row_number - 1 > settings.max_dataset_rows:
                raise HTTPException(
                    status_code=413,
                    detail={"code": "dataset_too_large", "message": f"Dataset exceeds the {settings.max_dataset_rows}-row MVP limit"},
                )
            if len(values) > len(columns) and any(value is not None for value in values[len(columns) :]):
                raise HTTPException(
                    status_code=400,
                    detail={"code": "extra_xlsx_fields", "message": f"XLSX row {row_number} has more fields than its header"},
                )
            row = list(values[: len(columns)])
            normalized_row: dict[str, Any] = {}
            for column, value in zip(columns, row):
                if isinstance(value, (datetime, date)):
                    normalized_row[column] = value.isoformat()
                else:
                    normalized_row[column] = value
            rows.append(normalized_row)
        workbook.close()
        return columns, rows
    except Exception as exc:
        if isinstance(exc, HTTPException):
            raise
        raise HTTPException(
            status_code=400,
            detail={"code": "invalid_xlsx", "message": "XLSX could not be parsed"},
        ) from exc


@router.post(
    "/upload",
    response_model=DatasetUploadResponse,
    status_code=status.HTTP_201_CREATED,
)
async def upload_dataset(
    file: UploadFile = File(...),
    user_id=Depends(current_user_id),
) -> DatasetUploadResponse:
    filename = Path(file.filename or "").name
    suffix = Path(filename).suffix.lower()
    if suffix not in {".csv", ".xlsx"}:
        raise HTTPException(
            status_code=400,
            detail={"code": "unsupported_file_type", "message": "Only CSV and XLSX are supported"},
        )

    content = await file.read(settings.max_upload_bytes + 1)
    if not content:
        raise HTTPException(
            status_code=400,
            detail={"code": "empty_file", "message": "Uploaded file is empty"},
        )
    if len(content) > settings.max_upload_bytes:
        raise HTTPException(
            status_code=413,
            detail={"code": "file_too_large", "message": "Uploaded file exceeds the configured limit"},
        )

    columns, rows = _parse_csv(content) if suffix == ".csv" else _parse_xlsx(content)
    if not rows:
        raise HTTPException(
            status_code=400,
            detail={"code": "empty_dataset", "message": "Dataset must contain at least one row"},
        )
    if len(rows) > settings.max_dataset_rows:
        raise HTTPException(
            status_code=413,
            detail={"code": "dataset_too_large", "message": f"Dataset exceeds the {settings.max_dataset_rows}-row MVP limit"},
        )

    dataset = repository.create_dataset(
        user_id=user_id,
        filename=filename,
        file_type=suffix[1:],
        columns=columns,
        rows=rows,
        preview=rows[:5],
    )
    return DatasetUploadResponse(
        dataset_id=dataset.id,
        filename=dataset.filename,
        file_type=dataset.file_type,
        row_count=len(dataset.rows),
        column_count=len(dataset.columns),
        columns=dataset.columns,
        preview=dataset.preview,
    )
